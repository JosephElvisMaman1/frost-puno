from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
OUTPUT = DOCS / "guion_exposicion_frost_puno.docx"


BLUE = RGBColor(11, 37, 69)
MID_BLUE = RGBColor(31, 77, 120)
ACCENT = RGBColor(46, 116, 181)
GRAY = RGBColor(89, 89, 89)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9E2EC")


def set_run(run, size: float = 10.5, color: RGBColor | None = None, bold: bool = False) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.text = "FrostPuno - Guion de exposicion"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run(run, size=9, color=GRAY)

    footer = section.footer.paragraphs[0]
    footer.text = "Material de apoyo para presentacion academica"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_run(run, size=9, color=GRAY)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, MID_BLUE),
        ("Heading 3", 12, ACCENT),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Guion de Exposicion - FrostPuno")
    set_run(run, size=22, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "App movil, despliegue y modelo ML v0.2.0 sin data leakage"
    )
    set_run(run, size=13, color=MID_BLUE)

    meta = [
        "Proyecto: Sistema inteligente distribuido para prediccion de heladas en Puno",
        "Curso: Aprendizaje de Maquina / Computacion Paralela y Distribuida",
        "Fecha de actualizacion: 2026-05-29",
    ]
    for line in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(line), size=10.5, color=GRAY)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        set_run(p.add_run(item), size=10.5)


def add_script_block(doc: Document, speaker: str, minutes: str, title: str, lines: list[str]) -> None:
    doc.add_heading(f"{speaker} - {title} ({minutes})", level=2)
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        set_run(p.add_run(line), size=10.5)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = [int(9360 / len(headers)) for _ in headers]
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_width(cell, widths[idx])
        set_cell_shading(cell, "E8EEF5")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run(run, size=9.5, color=BLUE, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            set_cell_width(cells[idx], widths[idx])
            for paragraph in cells[idx].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run(run, size=9)
    doc.add_paragraph()


def build_docx() -> None:
    doc = Document()
    style_document(doc)
    add_title(doc)

    doc.add_heading("1. Objetivo del guion", level=1)
    doc.add_paragraph(
        "Este guion esta preparado para una exposicion de tres integrantes. "
        "La idea es mostrar que FrostPuno evoluciono desde un prototipo tecnico hacia una app movil usable, desplegada y con un ciclo de mejora del modelo documentado."
    )

    doc.add_heading("2. Distribucion recomendada", level=1)
    add_table(
        doc,
        ["Integrante", "Tema principal", "Tiempo", "Idea que debe quedar clara"],
        [
            [
                "Expositor 1",
                "Problema, usuario y app movil",
                "4-5 min",
                "El usuario ya no llena un formulario tecnico; usa GPS, clima automatico y prediccion visual.",
            ],
            [
                "Expositor 2",
                "Arquitectura y despliegue",
                "4-5 min",
                "Flutter, FastAPI, Supabase, Render, Vercel y GitHub Actions estan separados y verificables.",
            ],
            [
                "Expositor 3",
                "Codigo y Machine Learning",
                "5-6 min",
                "v0.2.0 elimina data leakage, usa split por distrito y deja produccion intacta.",
            ],
        ],
    )

    doc.add_heading("3. Apertura breve", level=1)
    add_bullets(
        doc,
        [
            "Buenos dias. Presentamos FrostPuno, un sistema inteligente para estimar riesgo de heladas en la region Puno.",
            "El proyecto combina una app Flutter, un backend FastAPI, un pipeline de Machine Learning, Supabase y despliegue en la nube.",
            "La mejora principal es que la app ya no exige variables tecnicas al usuario: usa ubicacion, obtiene clima actual y genera una prediccion clara.",
        ],
    )

    add_script_block(
        doc,
        "Expositor 1",
        "4-5 min",
        "Problema y experiencia movil",
        [
            "El problema que abordamos son las heladas en zonas altoandinas de Puno, que afectan cultivos, ganado y la produccion tradicional de chuno.",
            "Nuestro objetivo no es reemplazar alertas oficiales, sino crear una base tecnologica academica que acerque informacion climatica y prediccion a usuarios no tecnicos.",
            "Al inicio la aplicacion parecia un formulario tecnico. El usuario tenia que ingresar temperatura, humedad, viento o nubosidad. Eso no era realista para una app movil.",
            "La interfaz actual se simplifico: el usuario presiona 'Usar mi ubicacion', la app obtiene latitud y longitud, consulta el backend en /weather/current y llena internamente las variables.",
            "Tambien se agrego manejo de permisos, fallback manual si falla GPS, modo oscuro, historial con cards y una pantalla de resultado mas visual.",
            "En el resultado se muestra el riesgo grande, por color, con confianza e icono. Esto hace que la decision sea mas rapida y entendible.",
        ],
    )

    add_script_block(
        doc,
        "Expositor 2",
        "4-5 min",
        "Arquitectura, Supabase y despliegue",
        [
            "La arquitectura esta separada por responsabilidades. Flutter solo consume la API; FastAPI concentra la logica de clima, prediccion, historial y metadata del modelo.",
            "El endpoint /weather/current consulta Open-Meteo por coordenadas. El endpoint /predict usa el modelo cargado desde el registry.",
            "Supabase queda preparado para persistencia de predicciones, feedback de usuarios y observaciones oficiales futuras. Se agregaron migraciones y politicas RLS.",
            "En despliegue, la API esta en Render en https://frost-puno.onrender.com y el frontend web en Vercel en https://frost-puno.vercel.app.",
            "GitHub Actions valida backend, Flutter, datos y modelo. Esto es importante porque el proyecto no depende solo de que funcione en mi computadora.",
            "Tambien existe render.v2.yaml para desplegar una API experimental con el modelo v0.2.0, sin romper la version productiva.",
        ],
    )

    add_script_block(
        doc,
        "Expositor 3",
        "5-6 min",
        "Codigo y aprendizaje de maquina",
        [
            "En el codigo, el pipeline ML vive en ml_pipeline. Ahi estan configuracion, ingesta, features, entrenamiento, evaluacion y registry.",
            "La version productiva v0.1.0 sigue en ml_pipeline/registry/frost_risk_model.joblib. Esa es la que usa FastAPI actualmente.",
            "Detectamos que las metricas perfectas de v0.1.0 eran optimistas por data leakage. El modelo recibia variables como temperatura_minima_diaria y horas_bajo_cero, que estaban muy relacionadas con la regla que generaba la etiqueta.",
            "Para corregirlo creamos v0.2.0. El dataset nuevo esta en data/processed/frost_training_dataset_v2.csv y elimina esas variables filtradas.",
            "v0.2.0 usa variables disponibles al momento de prediccion: temperatura actual, humedad, sensacion termica, punto de rocio, viento, nubosidad, precipitacion, latitud, longitud, altitud, mes y hora.",
            "Tambien cambiamos la evaluacion: usamos split por distrito, de modo que los distritos de test no aparecen en entrenamiento. Esto mide mejor generalizacion territorial.",
            "Las metricas bajaron: f1 macro queda cerca de 0.479. Eso no es un fracaso; es una evaluacion mas honesta que evita presentar un modelo artificialmente perfecto.",
            "La mejora continua no significa que el sistema cambie solo en produccion. Significa que se guardan feedback y observaciones, se reentrena una nueva version, se compara y se promueve manualmente si cumple criterios.",
        ],
    )

    doc.add_heading("4. Transiciones entre expositores", level=1)
    add_bullets(
        doc,
        [
            "Del expositor 1 al 2: 'Ahora que ya vimos la experiencia del usuario, pasamos a como esta construida y desplegada la arquitectura'.",
            "Del expositor 2 al 3: 'Con la arquitectura funcionando, la parte critica es el modelo: como se entrena, como se evalua y por que corregimos sus metricas'.",
            "Cierre del expositor 3: 'FrostPuno queda como un MVP funcional y extensible, con produccion estable y una ruta clara de mejora supervisada'.",
        ],
    )

    doc.add_heading("5. Preguntas probables y respuestas cortas", level=1)
    add_table(
        doc,
        ["Pregunta", "Respuesta recomendada"],
        [
            [
                "Por que v0.1.0 tenia F1 macro 1.00?",
                "Porque usaba variables derivadas de la regla de etiquetado. Eso es data leakage y hace que la evaluacion sea optimista.",
            ],
            [
                "v0.2.0 es peor porque baja metricas?",
                "No. Es mas realista porque elimina leakage y evalua distritos no vistos.",
            ],
            [
                "SENAMHI ya esta integrado oficialmente?",
                "Aun no como integracion completa. Hay una muestra curada y el pipeline queda preparado para validacion oficial progresiva.",
            ],
            [
                "El modelo aprende solo?",
                "No automaticamente en produccion. El ciclo es supervisado: datos, evaluacion, nueva version y promocion manual.",
            ],
            [
                "Que no se rompio?",
                "FastAPI, Supabase, Flutter, Render, Vercel y el modelo productivo v0.1.0 permanecen compatibles.",
            ],
        ],
    )

    doc.add_heading("6. Checklist antes de exponer", level=1)
    add_bullets(
        doc,
        [
            "Abrir https://frost-puno.vercel.app para mostrar frontend web.",
            "Abrir https://frost-puno.onrender.com/health para mostrar API viva.",
            "Tener capturas de GitHub, Render, Vercel y Supabase en el informe.",
            "Mostrar el APK o al menos la ruta del APK generado.",
            "Practicar la explicacion de data leakage con una frase simple: 'el modelo estaba viendo pistas de la respuesta'.",
        ],
    )

    doc.add_heading("7. Cierre sugerido", level=1)
    doc.add_paragraph(
        "FrostPuno demuestra una solucion academica completa: app movil usable, backend modular, datos abiertos, Supabase preparado, despliegue real y un pipeline ML versionado. "
        "La mejora mas importante no fue solo subir metricas, sino hacerlas mas confiables mediante una evaluacion sin data leakage."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_docx()
    print(OUTPUT)
