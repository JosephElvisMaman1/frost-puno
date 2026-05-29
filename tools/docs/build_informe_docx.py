from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
SOURCE = DOCS / "informe_frost_puno.md"
OUTPUT = DOCS / "informe_frost_puno_completo.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9E2EC")


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color in [
        ("Heading 1", 17, "0B2545"),
        ("Heading 2", 14, "1F4D78"),
        ("Heading 3", 12, "2E74B5"),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("FrostPuno")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(
        "Sistema inteligente distribuido para prediccion de heladas y apoyo a la produccion de chuno"
    )
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = RGBColor.from_string("1F4D78")

    meta = [
        "Universidad Nacional del Altiplano",
        "Facultad de Ingenieria Mecanica Electrica, Electronica y Sistemas",
        "Escuela Profesional de Ingenieria de Sistemas",
        "Autor: Joseph Elvis Mamani Mendoza",
        "Lugar: Puno, Peru",
        "Ano: 2026",
    ]
    for line in meta:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(line)

    callout = doc.add_paragraph()
    callout.alignment = WD_ALIGN_PARAGRAPH.CENTER
    callout.paragraph_format.space_before = Pt(18)
    run = callout.add_run(
        "Informe completo con arquitectura, ML v0.1.0/v0.2.0, despliegue, Supabase, GitHub Actions, Android APK y capturas."
    )
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("7A5A00")

    doc.add_section(WD_SECTION.NEW_PAGE)


def flush_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    max_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    width = int(9360 / max_cols)
    for row_idx, row in enumerate(rows):
        for col_idx in range(max_cols):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, width)
            text = row[col_idx] if col_idx < len(row) else ""
            cell.text = text.strip()
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
                    if row_idx == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor.from_string("0B2545")
            if row_idx == 0:
                set_cell_shading(cell, "E8EEF5")
    doc.add_paragraph()


def add_code_line(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string("0B2545")


def add_image(doc: Document, alt: str, target: str) -> None:
    image_path = (DOCS / target).resolve()
    if not image_path.exists():
        doc.add_paragraph(f"[Imagen no encontrada: {target}]")
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.2))
    caption = doc.add_paragraph(alt)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(9)


def clean_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("**", "")
    text = text.replace("*", "")
    return text.strip()


def build_docx() -> None:
    doc = Document()
    style_document(doc)
    add_cover(doc)

    in_code = False
    table_rows: list[list[str]] = []

    for raw_line in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw_line.rstrip()

        if table_rows and (not line.startswith("|") or re.match(r"^\|[\s:\-]+\|", line)):
            flush_table(doc, table_rows)
            table_rows = []
            if re.match(r"^\|[\s:\-]+\|", line):
                continue

        if line.startswith("```"):
            in_code = not in_code
            continue

        if in_code:
            add_code_line(doc, line)
            continue

        if not line.strip() or line.strip() == "---":
            if not table_rows:
                doc.add_paragraph()
            continue

        image = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if image:
            add_image(doc, clean_inline(image.group(1)), image.group(2))
            continue

        if line.startswith("|"):
            cells = [clean_inline(cell) for cell in line.strip("|").split("|")]
            if not all(re.fullmatch(r"[:\-\s]+", cell) for cell in cells):
                table_rows.append(cells)
            continue

        if line.startswith("# "):
            doc.add_heading(clean_inline(line[2:]), level=1)
            continue
        if line.startswith("## "):
            doc.add_heading(clean_inline(line[3:]), level=1)
            continue
        if line.startswith("### "):
            doc.add_heading(clean_inline(line[4:]), level=2)
            continue
        if line.startswith("#### "):
            doc.add_heading(clean_inline(line[5:]), level=3)
            continue

        if line.startswith("- "):
            doc.add_paragraph(clean_inline(line[2:]), style="List Bullet")
            continue

        if re.match(r"^\d+\.\s", line):
            doc.add_paragraph(clean_inline(re.sub(r"^\d+\.\s", "", line)), style="List Number")
            continue

        doc.add_paragraph(clean_inline(line))

    if table_rows:
        flush_table(doc, table_rows)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()
